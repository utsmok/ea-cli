from openpyxl import load_workbook
from pathlib import Path
root=Path(r'c:/dev/ea-cli/faculty_sheets')
target='20402048'
found_files=[]
for p in root.rglob('*.xlsx'):
    try:
        wb=load_workbook(p, data_only=True)
    except Exception:
        continue
    for sheet in wb.sheetnames:
        try:
            ws=wb[sheet]
            rows=list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers=[str(c).strip() if c is not None else '' for c in rows[0]]
            if 'material_id' not in headers:
                continue
            for i,row in enumerate(rows[1:], start=2):
                try:
                    mid=row[headers.index('material_id')]
                except Exception:
                    continue
                if str(mid).strip()==target:
                    print('FOUND in', p, 'sheet', sheet, 'row', i, row)
                    found_files.append((p, sheet, i, row))
        except Exception:
            continue
print('DONE. found', len(found_files))
