from pathlib import Path

import openpyxl
import polars as pl

from easy_access.sheets.sheet import store_complete_data


class DummyDataSettings:
    final_data_col_order = ["material_id", "a"]
    complete_data_name = "Complete Data"


class DummySettings:
    data_settings = DummyDataSettings()


def test_store_complete_data_writes_atomically(tmp_path: Path):
    settings = DummySettings()
    df = pl.DataFrame({"material_id": ["m1", "m2"], "a": [1, 2]})

    target = tmp_path / "out.xlsx"
    # ensure no leftover
    if target.exists():
        target.unlink()

    store_complete_data(settings, target, df)  # type: ignore

    assert target.exists(), "Target file must exist after store_complete_data"

    wb = openpyxl.load_workbook(filename=str(target))
    assert settings.data_settings.complete_data_name in wb.sheetnames
    sheet = wb[settings.data_settings.complete_data_name]
    # header + 2 rows
    assert sheet.max_row == df.shape[0] + 1

    # tmp file should not remain
    tmp_path_file = target.with_suffix(target.suffix + ".tmp")
    assert not tmp_path_file.exists()
